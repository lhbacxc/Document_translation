"""
将翻译稿 markdown 打包成 Word 文档（.docx）。

特性：
- 一级标题 `# xxx` -> Word Heading 1
- 二级标题 `## xxx` -> Word Heading 2
- 三级标题 `### xxx` -> Word Heading 3
- 普通段落 -> 正文（Times New Roman, 12 pt）
- Markdown 表格 -> Word 表格（自动识别和转换）
- 自动识别 `# References` 章节，使用 10 pt 字号
- 跳过空段
- 忽略 `>` 引用块（视为注释）

用法：
    python build_docx.py [输入md路径] [输出docx路径]

例：
    python build_docx.py 我的论文_translation_humanized.md
    # 输出到 output/我的论文_translation_humanized.docx

    python build_docx.py 我的论文_translation_humanized.md output/custom_name.docx
    # 自定义输出路径

默认：
    若不指定输入路径，默认读取 translation.md
    若不指定输出路径，自动基于输入文件名生成到 output/ 目录
"""

import sys
import os
import re

from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH


DEFAULT_FONT = "Times New Roman"
BODY_SIZE = Pt(12)
REF_SIZE = Pt(10)

SUPERSCRIPT_MAP = str.maketrans(
    "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ⁿᵟ",
    "0123456789+-=()nδ",
)
SUBSCRIPT_MAP = str.maketrans(
    "₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎",
    "0123456789+-=()",
)
_SUP_CHARS = {chr(k) for k in SUPERSCRIPT_MAP}
_SUB_CHARS = {chr(k) for k in SUBSCRIPT_MAP}


def parse_inline_scripts(text: str):
    """将文本拆分为 (片段, 'normal'|'super'|'sub') 列表。"""
    segments = []
    buf = []
    cur_type = "normal"

    def flush():
        if buf:
            segments.append(("".join(buf), cur_type))
            buf.clear()

    for ch in text:
        if ch in _SUP_CHARS:
            if cur_type != "super":
                flush()
                cur_type = "super"
            buf.append(ch.translate(SUPERSCRIPT_MAP))
        elif ch in _SUB_CHARS:
            if cur_type != "sub":
                flush()
                cur_type = "sub"
            buf.append(ch.translate(SUBSCRIPT_MAP))
        else:
            if cur_type != "normal":
                flush()
                cur_type = "normal"
            buf.append(ch)

    flush()
    return segments


def set_run_font(run, size: Pt):
    run.font.name = DEFAULT_FONT
    run.font.size = size
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        from docx.oxml import OxmlElement

        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), DEFAULT_FONT)
    rfonts.set(qn("w:hAnsi"), DEFAULT_FONT)
    rfonts.set(qn("w:eastAsia"), DEFAULT_FONT)


def _add_runs_with_scripts(container, text: str, size: Pt, bold: bool = False):
    """向 container（段落或标题）添加支持上下标的 run。"""
    for segment, seg_type in parse_inline_scripts(text):
        run = container.add_run(segment)
        set_run_font(run, size)
        if bold:
            run.bold = True
        if seg_type == "super":
            run.font.superscript = True
        elif seg_type == "sub":
            run.font.subscript = True


def add_heading(doc: Document, text: str, level: int):
    h = doc.add_heading(level=level)
    _add_runs_with_scripts(h, text, Pt(14 if level == 1 else 12), bold=True)


def add_paragraph(doc: Document, text: str, size: Pt = BODY_SIZE):
    p = doc.add_paragraph()
    _add_runs_with_scripts(p, text, size)


def add_table(doc: Document, rows: list, size: Pt = BODY_SIZE):
    """添加表格到 Word 文档。"""
    if not rows or len(rows) < 2:
        return

    # 创建表格
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = 'Light Grid Accent 1'

    # 填充表格内容
    for i, row_data in enumerate(rows):
        row_cells = table.rows[i].cells
        for j, cell_text in enumerate(row_data):
            if j < len(row_cells):
                cell = row_cells[j]
                # 清空单元格默认段落
                cell.text = ''
                # 添加支持上下标的内容
                p = cell.paragraphs[0]
                _add_runs_with_scripts(p, cell_text.strip(), size)
                # 表头加粗
                if i == 0:
                    for run in p.runs:
                        run.bold = True


def parse_markdown(md_text: str):
    """逐行解析 markdown，产出 (type, level, content, in_ref) 元组流。

    type ∈ {"heading", "paragraph", "table"}
    对于表格，content 是行列表 [[cell1, cell2, ...], ...]
    """
    blocks = []
    buf = []
    in_ref_section = False
    in_table = False
    table_rows = []

    def flush_buf():
        if buf:
            blocks.append(("paragraph", 0, " ".join(buf).strip(), in_ref_section))
            buf.clear()

    def flush_table():
        if table_rows:
            blocks.append(("table", 0, table_rows[:], in_ref_section))  # 使用 [:] 创建副本
            table_rows.clear()

    for raw in md_text.splitlines():
        line = raw.rstrip()

        # 跳过引用块
        if line.startswith(">"):
            continue

        # 检测标题
        m = re.match(r"^(#{1,6})\s+(.+)$", line)
        if m:
            flush_buf()
            flush_table()
            in_table = False
            level = len(m.group(1))
            title = m.group(2).strip()
            if title.lower().startswith("references"):
                in_ref_section = True
            else:
                if level == 1:
                    in_ref_section = False
            blocks.append(("heading", level, title, in_ref_section))
            continue

        # 检测表格行（必须在检测分隔线之前）
        if line.strip().startswith("|") and line.strip().endswith("|"):
            flush_buf()
            # 解析表格行
            cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
            # 跳过分隔符行（如 | --- | --- |）
            if all(re.match(r"^-+$", cell.strip()) for cell in cells if cell.strip()):
                in_table = True
                continue
            if in_table or not table_rows:
                table_rows.append(cells)
                in_table = True
            continue

        # 跳过 Markdown 水平分隔线（必须在表格检测之后）
        if line.strip().startswith("---") and not line.strip().startswith("|"):
            flush_buf()
            continue

        # 非表格行且之前在表格中，先 flush 表格
        if in_table and line.strip() and not line.strip().startswith("|"):
            flush_table()
            in_table = False

        # 空行
        if not line.strip():
            flush_buf()
            flush_table()
            in_table = False
            continue

        # 普通文本行
        buf.append(line.strip())

    flush_buf()
    flush_table()
    return blocks


def build_docx(md_path: str, out_path: str):
    with open(md_path, "r", encoding="utf-8") as f:
        md_text = f.read()

    blocks = parse_markdown(md_text)

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = DEFAULT_FONT
    style.font.size = BODY_SIZE

    for kind, level, content, in_ref in blocks:
        if kind == "heading":
            if not content:
                continue
            add_heading(doc, content, level=min(level, 4))
        elif kind == "paragraph":
            if not content:
                continue
            add_paragraph(doc, content, size=REF_SIZE if in_ref else BODY_SIZE)
        elif kind == "table":
            if content:
                add_table(doc, content, size=REF_SIZE if in_ref else BODY_SIZE)

    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    doc.save(out_path)
    print(f"[完成] 已生成 Word 文档 -> {out_path}")
    print(f"        共 {len(blocks)} 个段落/标题/表格块")


def main():
    md_path = sys.argv[1] if len(sys.argv) >= 2 else "translation.md"

    # 从输入 md 文件名推导输出 docx 文件名
    # 例如: 论文A_translation_humanized.md -> 论文A_translation_humanized.docx
    input_basename = os.path.splitext(os.path.basename(md_path))[0]
    default_out = os.path.join("output", f"{input_basename}.docx")

    out_path = sys.argv[2] if len(sys.argv) >= 3 else default_out

    if not os.path.isfile(md_path):
        print(f"[错误] 找不到翻译稿: {md_path}")
        sys.exit(1)

    build_docx(md_path, out_path)


if __name__ == "__main__":
    main()
